"""Migra los catalogos de ``Auxiliar.xlsx`` a la base de datos.

Ejecucion unica. Despues de esto la app deja de depender del Excel y los
catalogos se editan desde la pantalla de Ajustes.

Ademas de las columnas del Excel, siembra los valores que ya aparecen en los
registros existentes. Sin ese paso, un rig usado en una prueba vieja pero
ausente del Excel desapareceria del combo al editar ese registro.

    python tools/import_excel_catalogs.py
    python tools/import_excel_catalogs.py --dry-run
"""

from __future__ import annotations

import argparse
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from openpyxl import load_workbook  # noqa: E402

from app.config import PROJECT_ROOT, AppConfig  # noqa: E402
from app.db.connection import Database  # noqa: E402
from app.db.migrations import migrate  # noqa: E402
from app.models import EMPTY_RIG  # noqa: E402
from app.services.catalogs import color_for_index  # noqa: E402

EXCEL_PATH = PROJECT_ROOT / "excel_files" / "Auxiliar.xlsx"

# Columna del Excel -> (catalogo, tipo de prueba)
COLUMN_MAP = {
    "Cliente": ("customer", None),
    "Fatigue Rigs": ("rig", "fatigue"),
    "Torsion Rigs": ("rig", "torsion"),
    "Quasi Rigs": ("rig", "quasi"),
    "Fatigue Codes": ("code", "fatigue"),
    "Torsion Codes": ("code", "torsion"),
    "Quasi Codes": ("code", "quasi"),
    "Rotary Codes": ("code", "rotary"),
}

# El Excel no tiene columna de rigs de Rotary: el formulario los tenia escritos
# en el codigo, y ademas mal -- ``values="I-25"`` es una cadena, asi que el
# combo mostraba 'I', '-', '2', '5' como cuatro opciones sueltas.
ROTARY_RIG_SEED = ["I-25"]

# En el Excel, las celdas B11:B13 de la columna "Fatigue Rigs" no son rigs:
# son los tres resultados posibles de una muestra, guardados ahi por comodidad.
# La version anterior hacia df["Fatigue Rigs"].dropna().tolist() y por eso los
# ofrecia como si fueran bancos de prueba. Van a sample_statuses, no a rigs.
NOT_RIGS = {"FALLA", "SUSP", "S/FALLA"}

# Rigs por tabla de pruebas, para completar desde los datos existentes.
RIG_SOURCES = {
    "fatigue": ("fatigue_tests", [f"test_rig{i}" for i in range(1, 10)]),
    "torsion": ("torsion_tests", ["test_rig"]),
    "quasi": ("quasi_tests", ["test_rig"]),
    "rotary": ("rotary_tests", ["test_rig"]),
}

CUSTOMER_SOURCES = (
    "fatigue_tests", "rotary_tests", "torsion_tests", "quasi_tests",
)


def read_excel_columns() -> dict[str, list[str]]:
    """Lee las columnas del Excel sin pandas."""
    if not EXCEL_PATH.is_file():
        print(f"AVISO: no se encontro {EXCEL_PATH.name}, se usan solo los datos.")
        return {}

    workbook = load_workbook(EXCEL_PATH, read_only=True, data_only=True)
    sheet = workbook.active

    rows = list(sheet.iter_rows(values_only=True))
    workbook.close()

    if not rows:
        return {}

    headers = [str(h).strip() if h is not None else "" for h in rows[0]]
    columns: dict[str, list[str]] = {h: [] for h in headers if h}

    for row in rows[1:]:
        for header, value in zip(headers, row):
            if not header or value is None:
                continue
            text = str(value).strip()
            if text:
                columns[header].append(text)

    return columns


def values_from_data(db: Database, table: str, columns: list[str]) -> set[str]:
    """Valores distintos ya usados en los registros, para no perder ninguno."""
    if table not in db.table_names():
        return set()

    available = set(db.columns(table))
    found: set[str] = set()

    for column in columns:
        if column not in available:
            continue
        rows = db.query(
            f"SELECT DISTINCT {column} AS value FROM {table} "
            f"WHERE {column} IS NOT NULL AND TRIM({column}) != ''"
        )
        for row in rows:
            text = str(row["value"]).strip()
            if text and text != EMPTY_RIG:
                found.add(text)

    return found


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--dry-run", action="store_true",
        help="muestra lo que se importaria sin escribir",
    )
    parser.add_argument(
        "--database", metavar="RUTA",
        help="usa esta base en lugar de la de settings.json (para pruebas)",
    )
    args = parser.parse_args()

    config = AppConfig.load()
    target = Path(args.database) if args.database else config.database
    db = Database(target)

    if not db.exists():
        print(f"ERROR: no existe la base {target}")
        print("Revisa settings.json o corre primero tools/pick_database.py")
        return 1

    print(f"Base de datos: {target}")

    applied = migrate(db)
    if applied:
        print(f"Migraciones aplicadas: {', '.join(applied)}")

    excel = read_excel_columns()
    if excel:
        print(f"Columnas leidas de {EXCEL_PATH.name}: {', '.join(excel)}")

    # --- clientes --------------------------------------------------------
    customers: set[str] = set(excel.get("Cliente", []))
    for table in CUSTOMER_SOURCES:
        customers |= values_from_data(db, table, ["customer"])

    # --- rigs ------------------------------------------------------------
    rigs: dict[str, set[str]] = {}
    for column, (kind, test_type) in COLUMN_MAP.items():
        if kind == "rig":
            rigs.setdefault(test_type, set()).update(excel.get(column, []))

    rigs.setdefault("rotary", set()).update(ROTARY_RIG_SEED)

    for test_type, (table, columns) in RIG_SOURCES.items():
        rigs.setdefault(test_type, set()).update(
            values_from_data(db, table, columns)
        )

    # Quita los resultados de muestra que conviven con los rigs en el Excel.
    for test_type, names in rigs.items():
        rigs[test_type] = {n for n in names if n.upper() not in NOT_RIGS}

    # --- claves ----------------------------------------------------------
    codes: dict[str, set[str]] = {}
    for column, (kind, test_type) in COLUMN_MAP.items():
        if kind == "code":
            codes.setdefault(test_type, set()).update(
                value.upper() for value in excel.get(column, [])
            )

    # --- resumen ---------------------------------------------------------
    print(f"\nClientes: {len(customers)}")
    for test_type in sorted(rigs):
        print(f"Rigs {test_type}: {len(rigs[test_type])} -> "
              f"{', '.join(sorted(rigs[test_type])) or '(ninguno)'}")
    for test_type in sorted(codes):
        print(f"Claves {test_type}: {', '.join(sorted(codes[test_type])) or '(ninguna)'}")

    if args.dry_run:
        print("\n--dry-run: no se escribio nada.")
        return 0

    # --- escritura -------------------------------------------------------
    with db.write() as conn:
        for name in sorted(customers):
            conn.execute(
                "INSERT OR IGNORE INTO customers (name, active) VALUES (?, 1)",
                (name,),
            )

        color_index = 0
        for test_type in sorted(rigs):
            for name in sorted(rigs[test_type]):
                conn.execute(
                    "INSERT OR IGNORE INTO rigs (name, test_type, color, active) "
                    "VALUES (?, ?, ?, 1)",
                    (name, test_type, color_for_index(color_index)),
                )
                color_index += 1

        for test_type in sorted(codes):
            for code in sorted(codes[test_type]):
                conn.execute(
                    "INSERT OR IGNORE INTO test_codes (code, test_type) "
                    "VALUES (?, ?)",
                    (code, test_type),
                )

    totals = {
        "customers": db.query_one("SELECT COUNT(*) AS n FROM customers")["n"],
        "rigs": db.query_one("SELECT COUNT(*) AS n FROM rigs")["n"],
        "test_codes": db.query_one("SELECT COUNT(*) AS n FROM test_codes")["n"],
    }
    print("\nCatalogos en la base:")
    for name, count in totals.items():
        print(f"  {name:<12} {count}")

    print("\nListo. Los catalogos ya se editan desde Ajustes en la app.")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())

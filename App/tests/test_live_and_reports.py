"""Informacion al dia: datos nuevos de otros equipos y reportes a Excel.

1. Una pantalla abierta no se enteraba de lo que guardaba otra computadora
   hasta navegar o pulsar F5. Ahora la ventana pregunta al historial cada
   medio minuto y avisa -- sin recargar sola -- de lo que guardo otro equipo
   en lo que esa pantalla ensenia.
2. Solo se exportaban las fatigas en curso. Ahora tambien las finalizadas
   (con su periodo), Rotary, Torsion, Quasi, las Work Orders y el historial de
   mantenimiento.
"""

from harness import Report, make_context, offscreen, qt_app

offscreen()

import tempfile
from datetime import date, datetime, timedelta
from pathlib import Path

from openpyxl import load_workbook
from PySide6.QtWidgets import QFileDialog, QMessageBox

from app.models import (
    EMPTY_RIG,
    ONGOING,
    QUASI,
    SAMPLE_SLOTS,
    TORSION,
    WO_STARTED,
    AuditEntry,
    FatigueSample,
    FatigueTest,
    GenericTest,
    RigMaintenance,
    RotarySample,
    RotaryTest,
    WorkOrder,
)
from app.services import excel_export, maintenance
from app.services.identity import current_author
from app.ui.dialogs.maintenance_history_dialog import MaintenanceHistoryDialog
from app.ui.main_window import CHANGES_POLL_MS, MainWindow
from app.ui.widgets.changes_bar import describe_changes

# Lo propio lleva el equipo de verdad: es lo que la ventana descarta.
YO = current_author()
OTRO = ("luis", "PC-LAB-2")

report = Report("Información al día")
app = qt_app()
context = make_context()
carpeta = Path(tempfile.mkdtemp(prefix="bitacora_reportes_"))

avisos: list = []


def _sin_errores(*args, **_kwargs):
    raise AssertionError(f"se abrio un cuadro de error: {args[1:3]}")


QMessageBox.information = lambda *args, **_kwargs: avisos.append(args[1:3])
QMessageBox.warning = _sin_errores
QMessageBox.critical = _sin_errores

_consecutivo = [0]


def lote(tipo: str) -> str:
    _consecutivo[0] += 1
    return f"{995000 + _consecutivo[0]}{context.catalogs.codes(tipo)[0]}01"


def fatiga(autor, **cambios) -> int:
    prueba = FatigueTest(
        test_batch=lote("fatigue"), customer="AUDI",
        start_date=date(2026, 8, 1), qty_samples=1, test_status=ONGOING,
        samples=[FatigueSample(rig=context.catalogs.rig_names("fatigue")[0],
                               cycles=100)]
                + [FatigueSample() for _ in range(SAMPLE_SLOTS - 1)],
    )
    for campo, valor in cambios.items():
        setattr(prueba, campo, valor)
    return context.fatigue.create(prueba, autor)


# ======================================================================
report.section("1. El aviso dice quién, en qué y cuándo")

ahora = datetime(2026, 9, 11, 10, 0)


def apunte(quien, lote_, minutos, registro=1, tabla="fatigue_tests"):
    return AuditEntry(
        table_name=tabla, record_id=registro, test_batch=lote_,
        action="updated", field="cycles1", old_value="1", new_value="2",
        changed_by=quien, machine="PC", changed_at=ahora - timedelta(minutes=minutos),
    )


frase = describe_changes([apunte("luis", "241166STF09", 2)], ahora)
report.check("un cambio: quién, en qué registro y hace cuánto",
             frase == "luis guardó cambios en 241166STF09 hace 2 min.", frase)
frase = describe_changes([apunte("luis", "A", 5, 1), apunte("ana", "B", 1, 2),
                          apunte("luis", "A", 3, 1)], ahora)
report.check("varias personas y registros se resumen en una frase",
             frase == "luis y ana guardaron cambios en 2 registros hace 1 min.",
             frase)
report.check("lo de hace segundos dice 'hace un momento'",
             "hace un momento" in describe_changes([apunte("luis", "A", 0)], ahora))
report.check("un reloj adelantado en el otro equipo no da tiempos negativos",
             "hace un momento" in describe_changes([apunte("luis", "A", -3)], ahora))
report.check("y lo de horas se cuenta en horas",
             "hace 2 horas" in describe_changes([apunte("luis", "A", 125)], ahora))


# ======================================================================
report.section("2. El historial sabe qué guardó otro equipo")

marca = context.audit.last_id()
propia_id = fatiga(YO)
suya = context.fatigue.get(propia_id)
suya.samples[0].cycles = 500
context.fatigue.update(suya, OTRO)
torsion_id = context.torsion.create(
    GenericTest(test_batch=lote("torsion"), customer="AUDI",
                test_date=date(2026, 9, 1), qty_samples=1), OTRO)

ajenos = context.audit.changes_since(marca, exclude_machine=YO[1])
report.check("da lo que guardaron otros equipos desde la marca",
             {propia_id, torsion_id} <= {e.record_id for e in ajenos}
             and all(e.machine != YO[1] for e in ajenos),
             f"{len(ajenos)} apuntes")
report.check("sin lo que guardó este equipo",
             not any(e.action == "created" and e.record_id == propia_id
                     for e in ajenos))
solo_fatiga = context.audit.changes_since(marca, YO[1], tables={"fatigue_tests"})
report.check("filtra por las tablas que muestra cada pantalla",
             bool(solo_fatiga)
             and all(e.table_name == "fatigue_tests" for e in solo_fatiga))
report.check("una pantalla sin tablas no recibe nada",
             context.audit.changes_since(marca, YO[1], tables=set()) == [])


# ======================================================================
report.section("3. La ventana avisa sin recargar sola")

window = MainWindow(context)
barra = window.changes_bar


def visible() -> bool:
    return barra.isVisibleTo(window)


report.check(f"pregunta cada {CHANGES_POLL_MS // 1000} s",
             window.changes_timer.isActive()
             and window.changes_timer.interval() == CHANGES_POLL_MS)

window.show_page("fatigue")
window.check_for_changes()
report.check("recién cargada la pantalla, no hay aviso", not visible())

cambio = context.fatigue.get(propia_id)
cambio.samples[0].cycles = 900
context.fatigue.update(cambio, OTRO)
pantalla = window.pages["fatigue"].ongoing
antes = next(t for t in pantalla._records if t.id == propia_id).samples[0].cycles
window.check_for_changes()
texto = barra.message.text()
report.check("otro equipo guarda en Fatiga: aparece el aviso", visible())
report.check("y dice quién y en qué registro",
             "luis" in texto and cambio.test_batch in texto, texto)
report.check("con el detalle al pasar el cursor",
             "Ciclos" in barra.toolTip() and "PC-LAB-2" in barra.toolTip(),
             barra.toolTip())
report.check("sin recargar la pantalla por su cuenta", antes == 500, str(antes))

barra.refresh_button.click()
despues = next(t for t in pantalla._records if t.id == propia_id).samples[0].cycles
report.check("'Actualizar' trae los datos y oculta el aviso",
             not visible() and despues == 900, f"{despues} ciclos")

mia = context.fatigue.get(propia_id)
mia.comments = "propia"
context.fatigue.update(mia, YO)
window.check_for_changes()
report.check("lo que guarda este mismo equipo no se avisa", not visible())

# El mantenimiento vacia bancos de Fatiga: tambien cambia lo que se ve ahi.
abiertos = maintenance.open_by_rig(context.maintenance.open_records())
libre = next(r for r in context.catalogs.rigs("fatigue")
             if (r.name, "fatigue") not in abiertos)
paro_id = context.maintenance.start(
    RigMaintenance(rig_name=libre.name, test_type="fatigue", rig_id=libre.id,
                   start_date=date.today(), created_by="luis"), [], OTRO)
window.check_for_changes()
report.check("un mantenimiento de otro equipo avisa en Fatiga", visible(),
             barra.message.text())
window.refresh_current()

registro_torsion = context.torsion.get(torsion_id)
registro_torsion.comments = "revisada"
context.torsion.update(registro_torsion, OTRO)
window.check_for_changes()
report.check("un cambio en Torsión no avisa en la pantalla de Fatiga",
             not visible())

window.show_page("torsion")
window.check_for_changes()
report.check("y al entrar a Torsión ya viene cargado: tampoco avisa",
             not visible())

registro_torsion.comments = "revisada otra vez"
context.torsion.update(registro_torsion, OTRO)
window.check_for_changes()
report.check("en Torsión sí avisa", visible())
window.pages["torsion"].reload()
report.check("recargar con F5 también oculta el aviso", not visible())

window.show_page("dashboard")
registro_torsion.comments = "desde el dashboard"
context.torsion.update(registro_torsion, OTRO)
window.check_for_changes()
report.check("el dashboard, que resume todo, avisa de cualquier bitácora",
             visible())

window.show_menu()
registro_torsion.comments = "desde el menú"
context.torsion.update(registro_torsion, OTRO)
window.check_for_changes()
report.check("en el menú no hay nada que actualizar: no avisa", not visible())


# ======================================================================
report.section("4. Los reportes nuevos")

colores = context.catalogs.colors()
banco = context.catalogs.rig_names("fatigue")[0]
cerrada_id = fatiga(YO)
context.fatigue.close(cerrada_id, date(2026, 8, 20), YO)
cerrada = context.fatigue.get(cerrada_id)

ruta = excel_export.export_fatigue(
    [cerrada], colores, carpeta / "finalizadas.xlsx", finished=True,
    period=(date(2026, 8, 1), date(2026, 8, 31)))
hoja = load_workbook(ruta).active
cab = [c.value for c in hoja[2]]
report.check("finalizadas: el título dice qué y de qué periodo",
             "finalizadas" in hoja["A1"].value
             and "del 01/08/2026 al 31/08/2026" in hoja["A1"].value,
             hoja["A1"].value)
report.check("con la fecha de fin junto a la de inicio",
             cab[cab.index("Inicio") + 1] == "Fin", str(cab[:8]))
report.check("y la fecha de cierre en su celda",
             hoja.cell(row=3, column=cab.index("Fin") + 1).value == "20/08/2026")
celda = hoja.cell(row=3, column=cab.index("Test Rig 1") + 1)
report.check("los bancos siguen en su color",
             (celda.fill.fgColor.rgb or "")[-6:] == colores[banco].lstrip("#").upper(),
             f"{celda.value} {celda.fill.fgColor.rgb}")

en_curso = excel_export.export_fatigue_ongoing(
    context.fatigue.list(ONGOING), colores, carpeta / "en_curso.xlsx")
hoja = load_workbook(en_curso).active
report.check("el de en curso sale como siempre: sin columna Fin",
             "Fin" not in [c.value for c in hoja[2]]
             and "en curso" in hoja["A1"].value)

rig_rotary = context.catalogs.rig_names("rotary")[0]
rotary_id = context.rotary.create(RotaryTest(
    test_batch=lote("rotary"), customer="AUDI", start_date=date(2026, 9, 1),
    qty_samples=2, test_rig=rig_rotary,
    samples=[RotarySample(revs=1000, status="Falla"), RotarySample(revs=250)]
            + [RotarySample() for _ in range(SAMPLE_SLOTS - 2)]), YO)
rotarys = context.rotary.list()
hoja = load_workbook(excel_export.export_rotary(
    rotarys, colores, carpeta / "rotary.xlsx")).active
cab = [c.value for c in hoja[2]]
fila = next(r for r in range(3, hoja.max_row + 1)
            if hoja.cell(row=r, column=1).value == rotary_id)
report.check("Rotary: una fila por prueba y su total de revoluciones",
             hoja.max_row == len(rotarys) + 3
             and hoja.cell(row=fila, column=cab.index("Total revs") + 1).value == 1250,
             f"{hoja.max_row} filas para {len(rotarys)} pruebas")
report.check("con su banco en color",
             (hoja.cell(row=fila, column=cab.index("Rotary Rig") + 1)
              .fill.fgColor.rgb or "")[-6:]
             == colores[rig_rotary].lstrip("#").upper())

for config, repositorio in ((TORSION, context.torsion), (QUASI, context.quasi)):
    pruebas = repositorio.list()
    hoja = load_workbook(excel_export.export_generic(
        pruebas, config.label, colores, carpeta / f"{config.key}.xlsx")).active
    totales = next(r for r in range(3, hoja.max_row + 1)
                   if hoja.cell(row=r, column=1).value == "TOTALES")
    report.check(f"{config.label}: todas sus pruebas y el total de piezas",
                 hoja.max_row == len(pruebas) + 3
                 and hoja.cell(row=totales, column=6).value
                 == sum(p.qty_samples for p in pruebas),
                 f"{hoja.max_row} filas para {len(pruebas)} pruebas")

orden_id = context.work_orders.create(WorkOrder(
    test_type="fatigue", test_batch=lote("fatigue"), customer="AUDI",
    qty_samples=1, requester="Ana"), YO)
prueba_orden = fatiga(YO, test_batch=context.work_orders.get(orden_id).test_batch)
context.work_orders.mark_started(orden_id, prueba_orden, YO)
ordenes = context.work_orders.list()
hoja = load_workbook(excel_export.export_work_orders(
    ordenes, carpeta / "ordenes.xlsx")).active
cab = [c.value for c in hoja[2]]
filas = {hoja.cell(row=r, column=cab.index("Test Batch") + 1).value: r
         for r in range(3, hoja.max_row)}
comenzada = filas[context.work_orders.get(orden_id).test_batch]
report.check("Work Orders: la comenzada dice a qué registro dio lugar",
             hoja.cell(row=comenzada, column=cab.index("Estado") + 1).value
             == WO_STARTED
             and hoja.cell(row=comenzada, column=cab.index("Registro") + 1).value
             == f"#{prueba_orden}")
pendientes = [hoja.cell(row=r, column=1).value for r in range(3, hoja.max_row)
              if hoja.cell(row=r, column=cab.index("Estado") + 1).value != WO_STARTED]
report.check("y las pendientes van numeradas 1, 2, 3 como en pantalla",
             pendientes == list(range(1, len(pendientes) + 1)), str(pendientes[:5]))

periodos = context.maintenance.list()
hoja = load_workbook(excel_export.export_maintenance(
    periodos, carpeta / "mantenimiento.xlsx")).active
cab = [c.value for c in hoja[2]]
abiertos = [r for r in range(3, hoja.max_row)
            if hoja.cell(row=r, column=cab.index("Fin") + 1).value == "en curso"]
report.check("mantenimiento: un periodo por fila, los abiertos 'en curso'",
             hoja.max_row == len(periodos) + 3
             and len(abiertos) == sum(1 for p in periodos if p.is_open),
             f"{len(abiertos)} abiertos")


# ======================================================================
report.section("5. Cada pantalla exporta lo que se ve")

destino = [carpeta / "vacio.xlsx"]
QFileDialog.getSaveFileName = lambda *args, **kwargs: (str(destino[0]), "Excel")

window.show_page("fatigue")
finalizadas = window.pages["fatigue"].finished
finalizadas.refresh()
report.check("Fatiga finalizadas ya ofrece exportar",
             finalizadas.export_button.isVisibleTo(finalizadas))
destino[0] = carpeta / "boton_finalizadas.xlsx"
finalizadas.export_button.click()
hoja = load_workbook(destino[0]).active if destino[0].is_file() else None
report.check("y el botón escribe las pruebas que se ven, con su fecha de fin",
             hoja is not None
             and hoja.max_row >= len(finalizadas.table.visible_records()) + 3
             and "Fin" in [c.value for c in hoja[2]])

for clave in ("rotary", "torsion", "quasi", "work_orders"):
    window.show_page(clave)
    pagina = window.pages[clave]
    destino[0] = carpeta / f"boton_{clave}.xlsx"
    pagina.export_button.click()
    report.check(f"{clave}: el botón 'Exportar a Excel' escribe el archivo",
                 destino[0].is_file(), destino[0].name)

torsion = window.pages["torsion"]
window.show_page("torsion")
torsion.filters.search.setText("ZZZZ-no-existe")
app.processEvents()
destino[0] = carpeta / "no_debe_existir.xlsx"
avisos.clear()
torsion.export_button.click()
report.check("sin registros a la vista no genera archivo y lo dice",
             not destino[0].exists() and avisos and avisos[0][0] == "Sin datos",
             str(avisos))

historial = MaintenanceHistoryDialog(context.maintenance.list())
destino[0] = carpeta / "boton_mantenimiento.xlsx"
historial.export_button.click()
report.check("el historial de mantenimiento también se exporta",
             destino[0].is_file() and historial.isVisible() is False)

raise SystemExit(report.finish())

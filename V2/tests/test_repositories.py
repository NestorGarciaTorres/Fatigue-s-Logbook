"""Verifica repositorios, validacion, filtros y exportacion.

Cubre los pasos 4, 5, 6, 8, 11 y 13 del plan, sobre la copia migrada.
"""

from harness import (
    Report, database_copy, empty_database, make_context, offscreen,
    qt_app, settle,
)


import sys
from datetime import date
from pathlib import Path


from app.db.connection import Database
from app.db.migrations import migrate
from app.db.repositories import (
    AuditRepository, CatalogRepository, FatigueRepository,
    GenericRepository, RotaryRepository,
)
from app.models import FINISHED, ONGOING
from app.services import filtering, validation
from app.services.catalogs import CatalogService
from app.services.excel_export import export_fatigue_ongoing

COPY = database_copy()

failures = []


def check(label, condition, detail=""):
    status = "OK  " if condition else "FALLO"
    print(f"{status} {label}" + (f"   {detail}" if detail else ""))
    if not condition:
        failures.append(label)


def main():
    db = Database(COPY)
    # La copia se migra igual que al arrancar la app: sin esto le faltan
    # las tablas que agregan las migraciones nuevas.
    migrate(db)
    audit = AuditRepository(db)
    catalogs = CatalogService(CatalogRepository(db))
    catalogs.load()

    fatigue = FatigueRepository(db, audit)
    rotary = RotaryRepository(db, audit)
    torsion = GenericRepository(db, audit, "torsion_tests")
    quasi = GenericRepository(db, audit, "quasi_tests")

    print("=== Lectura de repositorios ===")
    all_fatigue = fatigue.list()
    check("fatigue: 628 registros", len(all_fatigue) == 628,
          f"leidos {len(all_fatigue)}")
    check("rotary: 13 registros", len(rotary.list()) == 13)
    check("torsion: 400 registros", len(torsion.list()) == 400)
    check("quasi: 63 registros", len(quasi.list()) == 63)

    ongoing = fatigue.list(ONGOING)
    finished = fatigue.list(FINISHED)
    check("ongoing + finished == total",
          len(ongoing) + len(finished) == len(all_fatigue),
          f"{len(ongoing)} + {len(finished)}")

    print("\n=== Totales contra el SQL de la app anterior ===")
    # Consulta literal de fatigue_cmplt_logbook.calculate_total_cycles()
    row = db.query_one("""
        SELECT COALESCE(SUM(COALESCE(cycles1,0) + COALESCE(cycles2,0) +
        COALESCE(cycles3,0) + COALESCE(cycles4,0) + COALESCE(cycles5,0) +
        COALESCE(cycles6,0) + COALESCE(cycles7,0) + COALESCE(cycles8,0) +
        COALESCE(cycles9,0)), 0), COALESCE(SUM(qty_samples), 0)
        FROM fatigue_tests WHERE test_status = 'Finished'
    """)
    old_cycles, old_samples = row[0], row[1]
    new_cycles = sum(t.total_cycles for t in finished)
    new_samples = sum(t.qty_samples for t in finished)

    check("ciclos coinciden", old_cycles == new_cycles,
          f"viejo {old_cycles:,} / nuevo {new_cycles:,}")
    check("piezas coinciden", old_samples == new_samples,
          f"viejo {old_samples:,} / nuevo {new_samples:,}")

    print("\n=== Ida y vuelta de fechas ===")
    with_dates = [t for t in all_fatigue if t.start_date]
    check("fechas parseadas", len(with_dates) == len(all_fatigue),
          f"{len(with_dates)}/{len(all_fatigue)}")
    sample = all_fatigue[0]
    print(f"     ejemplo: {sample.test_batch}  inicio={sample.start_date}"
          f"  fin={sample.end_date}  ciclos={sample.total_cycles:,}")

    print("\n=== Validacion ===")
    codes = catalogs.codes("fatigue")
    try:
        validation.validate_test_batch("242314STF09", codes)
        check("batch valido aceptado", True)
    except validation.ValidationError as e:
        check("batch valido aceptado", False, str(e))

    for bad, why in [("242314XXX09", "clave inexistente"),
                     ("24231STF09", "faltan digitos"),
                     ("", "vacio")]:
        try:
            validation.validate_test_batch(bad, codes)
            check(f"batch invalido rechazado ({why})", False)
        except validation.ValidationError:
            check(f"batch invalido rechazado ({why})", True)

    existing = all_fatigue[0].test_batch
    check("duplicado detectado", fatigue.batch_exists(existing))
    check("duplicado se ignora al editarse a si mismo",
          not fatigue.batch_exists(existing, exclude_id=all_fatigue[0].id))

    print("\n=== Filtros ===")
    total = len(all_fatigue)
    f_customer = filtering.TestFilters(customer=all_fatigue[0].customer)
    by_customer = filtering.apply(all_fatigue, f_customer)
    sql_count = db.query_one(
        "SELECT COUNT(*) AS n FROM fatigue_tests WHERE customer = ?",
        (all_fatigue[0].customer,))["n"]
    check("filtro por cliente cuadra con SQL",
          len(by_customer) == sql_count,
          f"filtro {len(by_customer)} / SQL {sql_count}")

    f_range = filtering.TestFilters(
        date_from=date(2026, 1, 1), date_to=date(2026, 12, 31))
    in_range = filtering.apply(all_fatigue, f_range)
    check("filtro por rango de fechas devuelve un subconjunto",
          0 < len(in_range) < total, f"{len(in_range)} de {total}")

    f_search = filtering.TestFilters(search="STF")
    check("busqueda libre encuentra resultados",
          len(filtering.apply(all_fatigue, f_search)) > 0)

    check("filtro vacio devuelve todo",
          len(filtering.apply(all_fatigue, filtering.TestFilters())) == total)

    print("\n=== Exportacion a Excel ===")
    out = Path(COPY).parent / "reporte.xlsx"
    export_fatigue_ongoing(ongoing, catalogs.colors(), out)
    check("archivo generado", out.is_file(), f"{out.stat().st_size:,} bytes")

    from openpyxl import load_workbook
    wb = load_workbook(out)
    ws = wb.active
    # fila 1 titulo, fila 2 encabezados, N filas, 1 de totales
    check("filas correctas en el Excel",
          ws.max_row == len(ongoing) + 3,
          f"{ws.max_row} filas para {len(ongoing)} registros")

    colored = sum(
        1 for row in ws.iter_rows(min_row=3, max_row=ws.max_row - 1)
        for cell in row
        if cell.fill and cell.fill.fgColor.rgb not in (None, "00000000")
    )
    check("hay celdas con color de rig", colored > 0, f"{colored} celdas")
    wb.close()

    print("\n" + "=" * 62)
    if failures:
        print(f"FALLARON {len(failures)}: " + ", ".join(failures))
        return 1
    print("TODAS LAS COMPROBACIONES PASARON")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())

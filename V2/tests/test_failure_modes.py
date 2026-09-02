"""Verifica la captura del modo de falla, de la base al Excel.

Corre contra una COPIA de la base real, para poder escribir de verdad.
"""

from harness import (
    Report, database_copy, empty_database, make_context, offscreen,
    qt_app, settle,
)


import os
import sqlite3
import sys
from datetime import date
from pathlib import Path

os.environ["QT_QPA_PLATFORM"] = "offscreen"

COPY = database_copy()

from openpyxl import load_workbook
from PySide6.QtWidgets import QApplication

from app.config import AppConfig
from app.context import AppContext
from app.models import (
    ONGOING, FailureMode, FatigueSample, FatigueTest, RotarySample, RotaryTest,
)
from app.services import excel_export, filtering
from app.ui import theme
from app.ui.dialogs.fatigue_dialog import FatigueDialog
from app.ui.dialogs.rotary_dialog import RotaryDialog
from app.ui.models.table_models import FatigueTableModel, RotaryTableModel
from app.ui.pages.settings_page import FailureModesTab
from app.ui.widgets.sample_chips import CHIPS_ROLE

failures = []
AUTHOR = ("verificador", "PC-PRUEBA")


def check(label, condition, detail=""):
    print(("OK    " if condition else "FALLO ") + label +
          (f"   {detail}" if detail else ""))
    if not condition:
        failures.append(label)


app = QApplication(sys.argv)
theme.apply(app)
context = AppContext(AppConfig(database_path=str(COPY), auto_backup=False))
print("migraciones:", context.prepare() or "ninguna pendiente")

catalogs = context.catalogs
repository = context.catalog_repository

# ======================================================================
print("\n=== 1. Catalogo de modos de falla ===")
seed = catalogs.failure_modes()
print(f"      {seed}")
check("el catalogo nace con modos de arranque", len(seed) == 7, str(len(seed)))

catalogs.save_failure_mode(FailureMode(label="Pandeo"))
check("agregar un modo lo deja disponible", "Pandeo" in catalogs.failure_modes())

added = next(m for m in repository.failure_modes() if m.label == "Pandeo")
added.active = False
catalogs.save_failure_mode(added)
check("desactivar lo saca del desplegable",
      "Pandeo" not in catalogs.failure_modes())
check("pero sigue en el catalogo completo",
      "Pandeo" in [m.label for m in repository.failure_modes(active_only=False)])

added.active = True
catalogs.save_failure_mode(added)

# ======================================================================
print("\n=== 2. Fatiga: ida y vuelta a la base ===")
test = FatigueTest(
    test_batch="999901STF01",
    customer="AUDI",
    start_date=date(2026, 8, 1),
    qty_samples=3,
    comments="prueba de modo de falla",
    samples=[
        FatigueSample(rig="I-02-1", cycles=120_000, failure_mode="Fractura"),
        FatigueSample(rig="Falla", cycles=88_000, failure_mode="Fisura"),
        FatigueSample(rig="S/Falla", cycles=500_000),      # sin modo: no fallo
    ] + [FatigueSample() for _ in range(6)],
)
new_id = context.fatigue.create(test, AUTHOR)
stored = context.fatigue.get(new_id)

check("el modo se guarda y se relee",
      [s.failure_mode for s in stored.samples[:3]] == ["Fractura", "Fisura", ""],
      str([s.failure_mode for s in stored.samples[:3]]))
check("una pieza sin modo vuelve como cadena vacia, no como --",
      stored.samples[2].failure_mode == "",
      repr(stored.samples[2].failure_mode))
check("no se perdieron rig ni ciclos",
      stored.samples[0].rig == "I-02-1" and stored.samples[0].cycles == 120_000)

raw = sqlite3.connect(COPY)
value = raw.execute(
    "SELECT failure_mode3 FROM fatigue_tests WHERE id = ?", (new_id,)
).fetchone()[0]
check("en la base la ranura vacia queda NULL", value is None, repr(value))
raw.close()

# --- auditoria del cambio ---
stored.samples[0].failure_mode = "Desgaste"
context.fatigue.update(stored, AUTHOR)
entries = context.audit.for_record("fatigue_tests", new_id)
changed = [e for e in entries if e.field == "failure_mode1"]
check("cambiar el modo queda en el historial", len(changed) == 1,
      str([(e.field, e.old_value, e.new_value) for e in changed]))
if changed:
    check("el historial guarda el valor anterior y el nuevo",
          (changed[0].old_value, changed[0].new_value) == ("Fractura", "Desgaste"),
          f"{changed[0].old_value} -> {changed[0].new_value}")

# --- registros heredados ---
# Un registro que de verdad venga de antes de la migracion: desde que la app
# esta en uso hay pruebas con modo capturado, y tomar la primera de la lista
# hacia fallar esto sin que nada estuviera mal.
legacy = next(t for t in context.fatigue.list(ONGOING)
              if t.id != new_id and not any(s.failure_mode for s in t.samples))
check("un registro anterior a la migracion no trae modos",
      all(s.failure_mode == "" for s in legacy.samples),
      str([s.failure_mode for s in legacy.samples[:3]]))
context.fatigue.update(legacy, AUTHOR)
reread = context.fatigue.get(legacy.id)
check("reguardarlo no inventa modos",
      all(s.failure_mode == "" for s in reread.samples))
check("reguardarlo no altera sus rigs",
      [s.rig for s in reread.samples] == [s.rig for s in legacy.samples])

# ======================================================================
print("\n=== 3. Rotary: ida y vuelta ===")
rotary = RotaryTest(
    test_batch="999902SRF01",
    customer="AUDI",
    start_date=date(2026, 8, 2),
    qty_samples=2,
    test_rig="I-25",
    samples=[
        RotarySample(revs=45_000, status="Falla", failure_mode="Fuga"),
        RotarySample(revs=90_000, status="S/Falla"),
    ] + [RotarySample() for _ in range(7)],
)
rotary_id = context.rotary.create(rotary, AUTHOR)
back = context.rotary.get(rotary_id)
check("Rotary guarda y relee el modo",
      [s.failure_mode for s in back.samples[:2]] == ["Fuga", ""],
      str([s.failure_mode for s in back.samples[:2]]))

# ======================================================================
print("\n=== 4. Formularios ===")
dialog = FatigueDialog(context.fatigue, catalogs, context.audit)
check("Fatiga: hay un desplegable de modo por cada una de las 9 piezas",
      len(dialog.failure_modes) == 9, str(len(dialog.failure_modes)))
combo = dialog.failure_modes[0]
items = [combo.itemText(i) for i in range(combo.count())]
print(f"      {items}")
check("ofrece todos los modos activos",
      items == catalogs.failure_modes(), str(items))
check("es una lista cerrada", not combo.isEditable())
check("abre vacio en un registro nuevo", combo.currentIndex() == -1,
      str(combo.currentText()))

editing = FatigueDialog(context.fatigue, catalogs, context.audit,
                        test=context.fatigue.get(new_id))
check("al editar, carga el modo capturado",
      editing.failure_modes[0].currentText() == "Desgaste",
      editing.failure_modes[0].currentText())
check("al editar, la pieza sin modo queda vacia",
      editing.failure_modes[2].currentText() == "",
      repr(editing.failure_modes[2].currentText()))
collected = editing._collect()
check("guardar recoge el modo del desplegable",
      [s.failure_mode for s in collected.samples[:3]] == ["Desgaste", "Fisura", ""],
      str([s.failure_mode for s in collected.samples[:3]]))

readonly = FatigueDialog(context.fatigue, catalogs, context.audit,
                         test=context.fatigue.get(new_id), read_only=True)
check("en solo lectura el modo no se puede cambiar",
      not readonly.failure_modes[0].isEnabled())

rotary_dialog = RotaryDialog(context.rotary, catalogs, context.audit,
                             test=context.rotary.get(rotary_id))
check("Rotary: 9 desplegables de modo", len(rotary_dialog.failure_modes) == 9)
check("Rotary: carga el modo capturado",
      rotary_dialog.failure_modes[0].currentText() == "Fuga",
      rotary_dialog.failure_modes[0].currentText())

# --- un registro que solo trae modo, sin rig ni ciclos ---
only_mode = FatigueTest(
    test_batch="999903STF01", customer="AUDI",
    start_date=date(2026, 8, 3), qty_samples=1,
    samples=[FatigueSample(failure_mode="Aflojamiento")]
             + [FatigueSample() for _ in range(8)],
)
only_id = context.fatigue.create(only_mode, AUTHOR)
loaded = FatigueDialog(context.fatigue, catalogs, context.audit,
                       test=context.fatigue.get(only_id))
check("una pieza con solo modo no se dibuja como ranura vacia",
      loaded.failure_modes[0].currentText() == "Aflojamiento",
      repr(loaded.failure_modes[0].currentText()))

# ======================================================================
print("\n=== 5. Tabla de la bitacora ===")
records = context.fatigue.list(ONGOING)
full = FatigueTableModel(catalogs, show_end_date=False, compact=False)
full.set_records(records)
mode_headers = [h for h in full.headers if h.startswith("Modo falla")]
check("la vista completa agrega una columna de modo por muestra",
      len(mode_headers) == 9, str(mode_headers[:3]))

row = next(i for i, r in enumerate(records) if r.id == new_id)
header_index = full.headers.index("Modo falla 1")
shown = full.data(full.index(row, header_index))
check("la columna de modo muestra el valor guardado", shown == "Desgaste", repr(shown))
check("la columna de rig sigue en su lugar",
      full.data(full.index(row, full.headers.index("Test Rig 1"))) == "I-02-1")
check("la columna de ciclos sigue en su lugar",
      full.data(full.index(row, full.headers.index("Ciclos 1"))) == "120,000")
check("las columnas de cola no se corrieron",
      full.data(full.index(row, full.headers.index("Total ciclos"))) == "708,000",
      full.data(full.index(row, full.headers.index("Total ciclos"))))

compact = FatigueTableModel(catalogs, show_end_date=False, compact=True)
compact.set_records(records)
chips = compact.data(compact.index(row, compact.chips_column), CHIPS_ROLE)
check("el chip lleva el modo en su tooltip",
      "Desgaste" in chips[0].tooltip, chips[0].tooltip)
check("el chip de la pieza sin modo no inventa texto",
      chips[2].tooltip.count("\u00b7") == 1, chips[2].tooltip)

rotary_model = RotaryTableModel(catalogs, compact=False)
rotary_model.set_records(context.rotary.list())
rrow = next(i for i, r in enumerate(rotary_model.records()) if r.id == rotary_id)
check("Rotary: la vista completa tambien trae el modo",
      rotary_model.data(
          rotary_model.index(rrow, rotary_model.headers.index("Modo falla 1"))
      ) == "Fuga")
check("Rotary: revs y estatus no se corrieron",
      rotary_model.data(
          rotary_model.index(rrow, rotary_model.headers.index("Revs 1"))
      ) == "45,000"
      and rotary_model.data(
          rotary_model.index(rrow, rotary_model.headers.index("Estatus 1"))
      ) == "Falla")

# ======================================================================
print("\n=== 6. Busqueda ===")
found = filtering.apply(records, filtering.TestFilters(search="Desgaste"))
check("se puede buscar por modo de falla",
      [t.id for t in found] == [new_id], str([t.id for t in found]))
rotary_found = filtering.apply(
    context.rotary.list(), filtering.TestFilters(search="Fuga")
)
check("tambien en Rotary", [t.id for t in rotary_found] == [rotary_id])

# ======================================================================
print("\n=== 7. Exportacion a Excel ===")
destination = COPY.parent / "export.xlsx"
excel_export.export_fatigue_ongoing(records, catalogs.colors(), destination)
book = load_workbook(destination)
sheet = book.active
headers = [c.value for c in sheet[2]]
check("el Excel trae una columna de modo por muestra",
      len([h for h in headers if str(h).startswith("Modo falla")]) == 9)
check("el orden por muestra es rig, ciclos, modo",
      headers[6:9] == ["Test Rig 1", "Ciclos 1", "Modo falla 1"],
      str(headers[6:9]))

excel_row = next(
    r for r in range(3, sheet.max_row + 1)
    if sheet.cell(row=r, column=2).value == "999901STF01"
)
mode_column = headers.index("Modo falla 1") + 1
check("el modo aparece en la celda correcta",
      sheet.cell(row=excel_row, column=mode_column).value == "Desgaste",
      repr(sheet.cell(row=excel_row, column=mode_column).value))
check("el modo de la segunda pieza tambien",
      sheet.cell(row=excel_row,
                 column=headers.index("Modo falla 2") + 1).value == "Fisura")
# openpyxl escribe la cadena vacia y la relee como None: la celda queda en
# blanco, que es lo que se busca. Es como se comporta ya la columna de rig.
blank = sheet.cell(row=excel_row, column=headers.index("Modo falla 3") + 1).value
check("la pieza sin modo exporta en blanco", blank in (None, ""), repr(blank))

rig_cell = sheet.cell(row=excel_row, column=headers.index("Test Rig 1") + 1)
expected = catalogs.colors().get("I-02-1", "").lstrip("#").upper()
actual = (rig_cell.fill.fgColor.rgb or "")[-6:].upper()
check("el color del rig sigue cayendo en la columna del rig",
      rig_cell.value == "I-02-1" and actual == expected,
      f"celda={rig_cell.value} color={actual} esperado={expected}")
check("la columna de modo no se pinto con el color del rig",
      (sheet.cell(row=excel_row, column=mode_column).fill.fgColor.rgb or "")[-6:]
      .upper() != expected)

totals_row = sheet.max_row
check("la fila de totales sigue cuadrando",
      sheet.cell(row=totals_row,
                 column=headers.index("Total ciclos") + 1).value
      == sum(t.total_cycles for t in records),
      str(sheet.cell(row=totals_row,
                     column=headers.index("Total ciclos") + 1).value))

# ======================================================================
print("\n=== 8. Ajustes ===")
tab = FailureModesTab(context)
tab.refresh()
labels = [tab.table.item(r, 0).text() for r in range(tab.table.rowCount())]
check("Ajustes lista todos los modos, activos y no",
      set(labels) == {m.label for m in repository.failure_modes(active_only=False)},
      str(labels))
check("cuenta cuantas muestras usan cada modo",
      tab.table.item(labels.index("Desgaste"), 2).text() == "1",
      tab.table.item(labels.index("Desgaste"), 2).text())
check("un modo sin uso cuenta cero",
      tab.table.item(labels.index("Pandeo"), 2).text() == "0")
check("el conteo de uso ve Fatiga y Rotary",
      repository.failure_mode_usage("Fuga") == 1,
      str(repository.failure_mode_usage("Fuga")))

app.quit()
print("\n" + "=" * 66)
if failures:
    print(f"FALLARON {len(failures)}:")
    for f in failures:
        print(f"  - {f}")
    sys.exit(1)
print("TODAS LAS COMPROBACIONES PASARON")

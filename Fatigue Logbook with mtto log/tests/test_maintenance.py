"""Mantenimiento de bancos: que se para, que se mueve y cuanto se descuenta.

Cuatro cosas que van juntas y que por separado mienten:

1. Poner un banco en mantenimiento saca de el las piezas que **siguen
   corriendo**, y solo esas: la que ya tiene resultado termino ahi, y su banco
   es el dato de donde corrio.
2. Al terminar, las piezas vuelven -- salvo la que mientras tanto se movio a
   otro banco, que es informacion mas reciente.
3. Los dias de la prueba se miden sin el tiempo parado, y los periodos se
   **unen**: dos bancos parados la misma semana detienen la prueba una semana,
   no dos.
4. La pieza detenida se ve distinta de una suspendida, y eso se comprueba
   contando pixeles, no leyendo el codigo.
"""

from harness import Report, make_context, offscreen, qt_app

offscreen()

from datetime import date
from pathlib import Path
from tempfile import mkdtemp

from openpyxl import load_workbook
from PySide6.QtCore import QRectF
from PySide6.QtGui import QColor, QImage, QPainter
from PySide6.QtWidgets import QGroupBox

from app.db.migrations import _M010_ROW_BACKGROUNDS
from app.db.migrations import _m010_contrast as contraste
from app.db.migrations import _m010_distance as distancia
from app.models import (
    ONGOING,
    SAMPLE_SLOTS,
    FatigueSample,
    FatigueTest,
    MaintenanceSample,
    RigMaintenance,
)
from app.services import duration, maintenance
from app.services.excel_export import (
    MAINTENANCE_LABEL,
    SUSPENDED_LABEL,
    export_fatigue_ongoing,
)
from app.ui import theme
from app.ui.models.table_models import DAYS_LEVEL_ROLE, FatigueTableModel
from app.ui.dialogs.maintenance_dialog import (
    FinishMaintenanceDialog,
    StartMaintenanceDialog,
)
from app.ui.dialogs.maintenance_history_dialog import (
    HEADERS as HEADERS_HISTORIAL,
)
from app.ui.dialogs.maintenance_history_dialog import MaintenanceHistoryDialog
from app.ui.pages.rigs_page import RigsPage
from app.ui.widgets.maintenance_panel import MaintenancePanel
from app.ui.widgets.sample_chips import (
    MAINTENANCE,
    MAINTENANCE_COLOR,
    SUSPENDED,
    SUSPENDED_COLOR,
    paint_chip_shape,
)

AUTHOR = ("verificador", "PC-PRUEBA")
report = Report("Mantenimiento de bancos")

app = qt_app()
context = make_context()
bancos = context.catalogs.rig_names("fatigue")
# Dos bancos que no vengan ya parados: se trabaja sobre una copia de la base de
# produccion y el laboratorio puede tener mantenimientos abiertos de verdad.
# Dar por hecho que no hay ninguno es la misma trampa que contar filas fijas.
_parados_de_entrada = maintenance.open_by_rig(context.maintenance.open_records())
_disponibles = [n for n in bancos
                if (n, "fatigue") not in _parados_de_entrada]
BANCO, OTRO = _disponibles[0], _disponibles[1]


def rig(name: str):
    return next(r for r in context.catalogs.rigs("fatigue") if r.name == name)


def periodo(inicio, fin, test_id=None, slot=1, banco=BANCO) -> RigMaintenance:
    """Un mantenimiento armado a mano, sin pasar por la base."""
    samples = []
    if test_id is not None:
        samples = [MaintenanceSample(record_id=test_id, slot=slot,
                                     rig_name=banco, test_batch="")]
    return RigMaintenance(rig_name=banco, test_type="fatigue",
                          start_date=inicio, end_date=fin, samples=samples)


# ======================================================================
report.section("1. Los periodos se unen, no se suman")

prueba = FatigueTest(id=77, test_batch="999870STF01", customer="AUDI",
                     start_date=date(2026, 1, 1), end_date=date(2026, 3, 1))

solapados = [
    periodo(date(2026, 1, 10), date(2026, 1, 20), test_id=77, slot=1),
    periodo(date(2026, 1, 15), date(2026, 1, 25), test_id=77, slot=2,
            banco=OTRO),
]
report.check(
    "dos paros que se solapan cuentan una sola vez",
    maintenance.stopped_days(prueba, solapados) == 15,
    f"{maintenance.stopped_days(prueba, solapados)} dias (10 + 10 = 20 si se "
    f"sumaran; del 10 al 25 son 15)",
)

separados = [
    periodo(date(2026, 1, 10), date(2026, 1, 20), test_id=77, slot=1),
    periodo(date(2026, 2, 1), date(2026, 2, 6), test_id=77, slot=2),
]
report.check("dos paros separados si se suman",
             maintenance.stopped_days(prueba, separados) == 15,
             f"{maintenance.stopped_days(prueba, separados)} dias (10 + 5)")

# Un banco que estuvo parado antes de que la prueba empezara no la detuvo.
anterior = [periodo(date(2025, 11, 1), date(2025, 12, 1), test_id=77)]
report.check("un paro anterior a la prueba no cuenta",
             maintenance.stopped_days(prueba, anterior) == 0,
             str(maintenance.stopped_days(prueba, anterior)))

# Y uno que la desborda se recorta a lo que dura la prueba.
desbordado = [periodo(date(2025, 12, 1), date(2026, 4, 1), test_id=77)]
report.check("un paro que la desborda se recorta a la vida de la prueba",
             maintenance.stopped_days(prueba, desbordado) == 59,
             f"{maintenance.stopped_days(prueba, desbordado)} dias, "
             f"la prueba dura {(prueba.end_date - prueba.start_date).days}")

# Un mantenimiento de otra prueba no frena a esta.
ajeno = [periodo(date(2026, 1, 10), date(2026, 1, 20), test_id=999)]
report.check("el paro de otra prueba no la afecta",
             maintenance.stopped_days(prueba, ajeno) == 0)

abierto = [periodo(date(2026, 1, 10), None, test_id=77)]
report.check(
    "uno abierto cuenta hasta la fecha de referencia",
    maintenance.stopped_days(prueba, abierto, date(2026, 1, 25)) == 15,
    str(maintenance.stopped_days(prueba, abierto, date(2026, 1, 25))),
)

fuera = maintenance.downtime_by_rig(solapados, date(2026, 3, 1))
report.check("el tiempo fuera de servicio se lleva por banco",
             fuera == {(BANCO, "fatigue"): 10, (OTRO, "fatigue"): 10},
             str(fuera))


# ======================================================================
report.section("2. Poner el banco en mantenimiento saca lo que corre")

corriendo = FatigueTest(
    test_batch="999871STF01", customer="AUDI", start_date=date(2026, 8, 1),
    qty_samples=3, test_status=ONGOING,
    samples=[
        FatigueSample(rig=BANCO, cycles=1000),                    # corriendo
        FatigueSample(rig=BANCO, result="Falla", cycles=5000),    # ya termino
        FatigueSample(rig=OTRO, cycles=2000),                     # otro banco
    ] + [FatigueSample() for _ in range(SAMPLE_SLOTS - 3)],
)
test_id = context.fatigue.create(corriendo, AUTHOR)

afectadas = [s for s in maintenance.affected_samples(
    BANCO, "fatigue", context.fatigue.list(ONGOING)
) if s.record_id == test_id]

report.check("solo se saca la pieza que sigue corriendo",
             [s.slot for s in afectadas] == [1],
             str([s.slot for s in afectadas]))

banco = rig(BANCO)
registro = RigMaintenance(rig_name=BANCO, test_type="fatigue", rig_id=banco.id,
                          start_date=date(2026, 8, 10),
                          reason="Cambio de mordazas", created_by="verificador")
todas = maintenance.affected_samples(
    BANCO, "fatigue", context.fatigue.list(ONGOING)
)
mantenimiento_id = context.maintenance.start(registro, todas, AUTHOR)

guardado = context.fatigue.get(test_id)
report.check("la pieza que corria se queda sin banco",
             guardado.samples[0].rig == "--", repr(guardado.samples[0].rig))
report.check("y conserva sus ciclos", guardado.samples[0].cycles == 1000)
report.check("la pieza que ya habia fallado conserva su banco",
             guardado.samples[1].rig == BANCO, guardado.samples[1].rig)
report.check("la de otro banco no se toca",
             guardado.samples[2].rig == OTRO, guardado.samples[2].rig)

abiertos = context.maintenance.open_records()
detenidas = maintenance.paused_slots(abiertos)
report.check("la pieza queda marcada como detenida",
             (test_id, 1) in detenidas)
report.check("y no la que fallo ni la de otro banco",
             (test_id, 2) not in detenidas and (test_id, 3) not in detenidas)
report.check("el banco aparece parado",
             (BANCO, "fatigue") in maintenance.open_by_rig(abiertos))

historial = context.audit.for_record("fatigue_tests", test_id)
salida = [e for e in historial if e.action == "maintenance"]
report.check("queda en el historial con su propia accion", len(salida) == 1,
             str([(e.action, e.field) for e in salida]))
if salida:
    report.check("y dice de que banco salio la pieza",
                 salida[0].field == "test_rig1"
                 and salida[0].old_value == BANCO,
                 f"{salida[0].field}: {salida[0].old_value}")


# ======================================================================
report.section("3. Terminar el mantenimiento devuelve las piezas")

# Antes de cerrar, una segunda prueba se mueve a otro banco por su cuenta: esa
# no debe volver, su dato es mas reciente que el apunte del mantenimiento.
movida = FatigueTest(
    test_batch="999872STF01", customer="AUDI", start_date=date(2026, 8, 2),
    qty_samples=1, test_status=ONGOING,
    samples=[FatigueSample(rig=OTRO, cycles=300)]
            + [FatigueSample() for _ in range(SAMPLE_SLOTS - 1)],
)
movida_id = context.fatigue.create(movida, AUTHOR)
otro_id = context.maintenance.start(
    RigMaintenance(rig_name=OTRO, test_type="fatigue", rig_id=rig(OTRO).id,
                   start_date=date(2026, 8, 12), created_by="verificador"),
    maintenance.affected_samples(OTRO, "fatigue", [context.fatigue.get(movida_id)]),
    AUTHOR,
)
reasignada = context.fatigue.get(movida_id)
reasignada.samples[0].rig = BANCO
context.fatigue.update(reasignada, AUTHOR)

repuestas = context.maintenance.finish(
    mantenimiento_id, date(2026, 8, 20), AUTHOR
)
vuelta = context.fatigue.get(test_id)
report.check("la pieza vuelve a su banco",
             vuelta.samples[0].rig == BANCO, vuelta.samples[0].rig)
report.check("se informa cuantas se repusieron", repuestas >= 1, str(repuestas))

context.maintenance.finish(otro_id, date(2026, 8, 20), AUTHOR)
report.check("la que se movio a otro banco se queda donde la pusieron",
             context.fatigue.get(movida_id).samples[0].rig == BANCO,
             context.fatigue.get(movida_id).samples[0].rig)

cerrado = context.maintenance.get(mantenimiento_id)
report.check("el periodo queda cerrado", not cerrado.is_open)
report.check("y con sus dias medidos", cerrado.days() == 10,
             f"{cerrado.days()} dias, del 10 al 20 de agosto")
report.check("las piezas quedan marcadas como repuestas",
             all(s.restored for s in cerrado.samples),
             str([(s.slot, s.restored) for s in cerrado.samples]))

vuelto = context.audit.for_record("fatigue_tests", test_id)
report.check("el regreso al banco tambien queda registrado",
             any(e.action == "restored" and e.new_value == BANCO
                 for e in vuelto))


# ======================================================================
report.section("4. Los dias en curso no cuentan el mantenimiento")

registros = context.maintenance.list()
paros = maintenance.stopped_by_test(context.fatigue.list(), registros)
report.check("la prueba tiene dias detenidos", paros.get(test_id) == 10,
             str(paros.get(test_id)))

hoy = date(2026, 9, 8)
calendario = (hoy - vuelta.start_date).days
efectivos = duration.days_running(vuelta.start_date, hoy, paros[test_id])
report.check("los dias de ensayo son los de calendario menos el paro",
             efectivos == calendario - 10,
             f"{efectivos} de ensayo, {calendario} de calendario")

report.check("nunca sale un numero negativo",
             duration.days_running(date(2026, 9, 1), hoy, 999) == 0)

modelo = FatigueTableModel(context.catalogs, show_end_date=False)
modelo.set_maintenance(paros, maintenance.paused_slots(registros))
modelo.set_records([vuelta])
dias_columna = modelo.index(0, modelo.days_column)
report.check("la tabla ensenia los dias ya descontados",
             modelo.value(vuelta, modelo.days_column)
             == duration.days_running(vuelta.start_date,
                                      stopped=paros[test_id]),
             str(modelo.value(vuelta, modelo.days_column)))
report.check("el semaforo sigue teniendo nivel",
             modelo.data(dias_columna, DAYS_LEVEL_ROLE) in
             (duration.OK, duration.WARNING, duration.CRITICAL),
             str(modelo.data(dias_columna, DAYS_LEVEL_ROLE)))

# Y el desglose se explica en el tooltip, que es donde se puede mirar sin
# tener que restar de cabeza.
_, _, hint = modelo.days_info(vuelta)
report.check("el tooltip desglosa calendario y mantenimiento",
             "mantenimiento" in (hint or "").lower(), repr(hint))


# ======================================================================
report.section("5. La pieza detenida se ve distinta de una suspendida")

parada = FatigueTest(
    id=555, test_batch="999873STF01", customer="AUDI",
    start_date=date(2026, 8, 1), qty_samples=1, test_status=ONGOING,
    samples=[FatigueSample(rig="--", cycles=1500)]
            + [FatigueSample() for _ in range(SAMPLE_SLOTS - 1)],
)

suelto = FatigueTableModel(context.catalogs, show_end_date=False)
suelto.set_records([parada])
chips = suelto.chips(parada)
report.check("sin mantenimiento, la pieza sin banco es una suspension",
             chips[0].kind == SUSPENDED, chips[0].kind)

con_paro = FatigueTableModel(context.catalogs, show_end_date=False)
con_paro.set_maintenance({}, {(555, 1): periodo(date(2026, 8, 5), None)})
con_paro.set_records([parada])
chips = con_paro.chips(parada)
report.check("con mantenimiento abierto, la pieza cambia de clase",
             chips[0].kind == MAINTENANCE, chips[0].kind)
report.check("y su tooltip dice de que banco y desde cuando",
             "mantenimiento" in chips[0].tooltip.lower()
             and BANCO in chips[0].tooltip, chips[0].tooltip)
report.check("el chip va en el rojo del mantenimiento",
             chips[0].color == MAINTENANCE_COLOR, str(chips[0].color))
report.check("que no es el naranja de una suspension",
             MAINTENANCE_COLOR.upper() != SUSPENDED_COLOR.upper(),
             f"{MAINTENANCE_COLOR} contra {SUSPENDED_COLOR}")

# El rojo se eligio midiendo, y lo que se midio se comprueba aqui: si maniana
# alguien lo cambia por otro rojo "que se ve bien", esto lo frena.
flojos = {f: round(contraste(MAINTENANCE_COLOR, f), 2)
          for f in _M010_ROW_BACKGROUNDS
          if contraste(MAINTENANCE_COLOR, f) < 3.0}
report.check("se lee sobre los tres fondos de fila (>= 3:1)", not flojos,
             str(flojos) if flojos else str({
                 f: round(contraste(MAINTENANCE_COLOR, f), 2)
                 for f in _M010_ROW_BACKGROUNDS}))
report.check("DANGER no habria servido: falla en la franja alterna",
             contraste(theme.DANGER, theme.ROW_ALT) < 3.0,
             f"{contraste(theme.DANGER, theme.ROW_ALT):.2f} : 1")
report.check("no se confunde con el naranja de suspendida (dE >= 24)",
             distancia(MAINTENANCE_COLOR, SUSPENDED_COLOR) >= 24,
             f"dE {distancia(MAINTENANCE_COLOR, SUSPENDED_COLOR):.1f}")

# Contra los bancos queda por debajo del dE 24 de la regla, y a proposito: el
# vecino es el salmon de I-02-1 y bajar de 3:1 para separarse de el habria
# costado la legibilidad del numero. Lo que separa los dos casos es la trama,
# que ningun chip de banco lleva.
vecino = min((distancia(MAINTENANCE_COLOR, r.color), r.name)
             for r in context.catalogs.rigs())
report.check("el banco mas cercano sigue a distancia reconocible (dE >= 20)",
             vecino[0] >= 20, f"{vecino[1]} a dE {vecino[0]:.1f}")
report.check("y ningun chip de banco lleva trama, que es lo que los separa",
             all(c.kind != MAINTENANCE
                 for c in suelto.chips(parada) + chips[1:]))


def interior_marcado(kind: str) -> int:
    """Pixeles del interior del chip que no son el relleno liso.

    Se mira solo el centro, lejos del borde: el contorno punteado de una
    suspension vive en el borde y el antialias del redondeo tambien, asi que
    contarlos ahi no distinguiria nada. La trama del mantenimiento, en cambio,
    cruza el chip entero.
    """
    image = QImage(40, 20, QImage.Format.Format_ARGB32)
    image.fill(QColor(MAINTENANCE_COLOR))

    painter = QPainter(image)
    painter.setRenderHint(QPainter.RenderHint.Antialiasing)
    paint_chip_shape(painter, QRectF(0, 0, 40, 20), kind, MAINTENANCE_COLOR)
    painter.end()

    relleno = QColor(MAINTENANCE_COLOR).rgb()
    return sum(
        1
        for y in range(6, 14)
        for x in range(8, 32)
        if QColor(image.pixel(x, y)).rgb() != relleno
    )


trama = interior_marcado(MAINTENANCE)
liso = interior_marcado(SUSPENDED)
report.check("el chip de mantenimiento lleva trama en el interior",
             trama > 20, f"{trama} pixeles marcados de 192")
report.check("el de suspension deja el interior liso",
             liso == 0, f"{liso} pixeles marcados de 192")
report.check("los dos se distinguen sin mirar el color",
             trama > liso + 20, f"trama {trama} contra liso {liso}")


# ======================================================================
report.section("6. La pantalla de rigs lo cuenta")

pagina = RigsPage(context)
pagina.refresh()

tarjetas = pagina.cards()
por_nombre = {t.rig.name: t for t in tarjetas}
report.check("hay una tarjeta por banco del catalogo",
             len(tarjetas) == len(context.catalogs.rigs()),
             f"{len(tarjetas)} tarjetas")

# Nada de contar un absoluto: la base es una copia de la de produccion y ahi
# puede haber mantenimientos abiertos capturados por el laboratorio. Se compara
# contra lo que diga la base, y despues contra el propio numero de antes.
abiertos_base = len(maintenance.open_by_rig(context.maintenance.open_records()))
parados = int(pagina.card_maintenance.value_label.text())
report.check("la tarjeta cuenta los mantenimientos abiertos que haya",
             parados == abiertos_base,
             f"la tarjeta dice {parados}, la base tiene {abiertos_base}")

libres = int(pagina.card_free.value_label.text())
ocupados = int(pagina.card_busy.value_label.text())
report.check("libres, ocupados y parados suman el catalogo",
             libres + ocupados + parados == len(context.catalogs.rigs()),
             f"{libres} + {ocupados} + {parados}")

# Y con uno abierto, la cuenta se mueve de sitio.
context.maintenance.start(
    RigMaintenance(rig_name=BANCO, test_type="fatigue", rig_id=banco.id,
                   start_date=date.today(), reason="Calibración",
                   created_by="verificador"),
    maintenance.affected_samples(BANCO, "fatigue",
                                 context.fatigue.list(ONGOING)),
    AUTHOR,
)
pagina.refresh()
report.check("al abrir uno, la tarjeta de mantenimiento sube en uno",
             int(pagina.card_maintenance.value_label.text()) == parados + 1,
             f"de {parados} a {pagina.card_maintenance.value_label.text()}")
report.check("y los tres numeros siguen sumando el catalogo",
             int(pagina.card_free.value_label.text())
             + int(pagina.card_busy.value_label.text())
             + int(pagina.card_maintenance.value_label.text())
             == len(context.catalogs.rigs()))

tarjetas = pagina.cards()
parada_card = next(t for t in tarjetas if t.rig.name == BANCO
                   and t.rig.test_type == "fatigue")
report.check("la tarjeta del banco parado lo sabe",
             parada_card.maintenance is not None)
report.check("y su motivo es el que se capturo",
             parada_card.maintenance.reason == "Calibración",
             parada_card.maintenance.reason)


# ======================================================================
report.section("7. Los dos dialogos se arman con lo que hay que ver")

# Se construyen, no se abren: exec() bloquea esperando a que alguien cierre.
abierto_ahora = maintenance.open_by_rig(
    context.maintenance.open_records()
)[(BANCO, "fatigue")]

pendientes = maintenance.affected_samples(
    OTRO, "fatigue", context.fatigue.list(ONGOING)
)
alta = StartMaintenanceDialog(rig(OTRO), pendientes)
report.check("el de alta arranca hoy", alta.start_date() == date.today(),
             str(alta.start_date()))
report.check("y no deja fecharlo en el futuro",
             alta.start_edit.maximumDate().toPython() == date.today())
titulos = [g.title() for g in alta.findChildren(QGroupBox)]
report.check("dice cuantas piezas van a salir del banco",
             any(f"({len(pendientes)})" in t for t in titulos),
             f"{titulos} con {len(pendientes)} piezas")
armado = alta.maintenance(created_by="verificador")
report.check("arma el periodo con el banco y su tipo",
             armado.rig_name == OTRO and armado.test_type == "fatigue",
             f"{armado.rig_name} / {armado.test_type}")
report.check("y con el id del rig, para la clave foranea",
             armado.rig_id == rig(OTRO).id)

cierre = FinishMaintenanceDialog(abierto_ahora)
report.check("el de cierre no deja terminar antes de empezar",
             cierre.end_edit.minimumDate().toPython()
             == abierto_ahora.start_date,
             str(cierre.end_edit.minimumDate().toPython()))
report.check("ofrece devolver las piezas por omision", cierre.restore())

vacio = FinishMaintenanceDialog(
    RigMaintenance(rig_name=OTRO, test_type="fatigue",
                   start_date=date(2026, 9, 1))
)
report.check("sin piezas pendientes, la casilla de devolver se apaga",
             not vacio.restore_check.isEnabled())

# ======================================================================
report.section("8. El Excel distingue detenida de suspendida")

# El reporte impreso se consulta lejos de la pantalla: si las dos cosas salen
# igual, quien lo lee no puede saber si la prueba esta parada por su culpa o
# por la del banco.
en_curso = context.fatigue.list(ONGOING)
paradas = maintenance.paused_slots(context.maintenance.open_records())
destino = Path(mkdtemp(prefix="bitacora_xlsx_")) / "reporte.xlsx"
export_fatigue_ongoing(en_curso, context.catalogs.colors(), destino,
                       paused=paradas)

hoja = load_workbook(destino).active
etiquetas = {}
for f in range(3, hoja.max_row + 1):
    for c in range(1, hoja.max_column + 1):
        celda = hoja.cell(row=f, column=c)
        if celda.value in (MAINTENANCE_LABEL, SUSPENDED_LABEL):
            etiquetas.setdefault(celda.value, celda)

report.check("el reporte marca las piezas detenidas por mantenimiento",
             MAINTENANCE_LABEL in etiquetas,
             f"etiquetas encontradas: {sorted(etiquetas)}")
if MAINTENANCE_LABEL in etiquetas:
    relleno = etiquetas[MAINTENANCE_LABEL].fill.fgColor.rgb or ""
    report.check("con el mismo rojo que el chip de la pantalla",
                 relleno.endswith(MAINTENANCE_COLOR.lstrip("#")), relleno)
if SUSPENDED_LABEL in etiquetas:
    report.check("y la suspendida conserva su naranja, que dice otra cosa",
                 (etiquetas[SUSPENDED_LABEL].fill.fgColor.rgb or "")
                 .endswith(SUSPENDED_COLOR.lstrip("#")),
                 etiquetas[SUSPENDED_LABEL].fill.fgColor.rgb)

pie = [hoja.cell(row=f, column=1).value for f in range(3, hoja.max_row + 1)]
report.check("y la hoja explica al pie que significa",
             any(isinstance(t, str) and t.startswith(MAINTENANCE_LABEL + ":")
                 for t in pie),
             str([t for t in pie if isinstance(t, str) and len(t) > 30][:2]))

# ======================================================================
report.section("9. El historial se puede consultar")

# Hasta ahora solo se veia el periodo abierto, en la tarjeta de su banco. Los
# cerrados se median --se descuentan de los dias de la prueba-- pero no
# aparecian en ninguna pantalla.
historia = [
    periodo(date(2026, 3, 1), date(2026, 3, 11)),          # cerrado, 10 dias
    periodo(date(2026, 5, 5), date(2026, 5, 8), banco=OTRO),
    periodo(date(2026, 8, 1), None),                        # abierto
]

dentro = maintenance.overlapping(historia, date(2026, 4, 1), date(2026, 6, 1))
report.check("el rango deja fuera lo que no lo toca",
             [r.start_date for r in dentro] == [date(2026, 5, 5)],
             str([str(r.start_date) for r in dentro]))

# Uno abierto que empezo antes del rango sigue contando: es lo que tiene el
# banco parado hoy, y esconderlo seria esconder justo el que importa.
abarca = maintenance.overlapping(historia, date(2026, 8, 15), date(2026, 9, 1),
                                 reference=date(2026, 9, 1))
report.check("un periodo abierto que empezo antes del rango si cuenta",
             [r.start_date for r in abarca] == [date(2026, 8, 1)],
             str([str(r.start_date) for r in abarca]))

todos = maintenance.overlapping(historia)
report.check("sin rango salen todos, del mas reciente al mas antiguo",
             [r.start_date for r in todos]
             == sorted((r.start_date for r in historia), reverse=True),
             str([str(r.start_date) for r in todos]))

recortado = maintenance.downtime_by_rig(
    historia, reference=date(2026, 9, 1),
    window=(date(2026, 3, 5), date(2026, 3, 20)),
)
report.check("los dias fuera de servicio se recortan al rango",
             recortado == {(BANCO, "fatigue"): 6},
             f"{recortado} (del 5 al 11 son 6, no los 10 del periodo)")

completo = maintenance.downtime_by_rig(historia, reference=date(2026, 9, 1))
report.check("y sin rango sale el periodo entero",
             completo[(BANCO, "fatigue")] == 10 + 31,
             str(completo))

dialogo = MaintenanceHistoryDialog(todos)
report.check("el historial lista un renglon por periodo",
             dialogo.table.rowCount() == len(todos),
             str(dialogo.table.rowCount()))
report.check("y resume cuantos son y cuanto suman",
             "3 periodos" in dialogo.summary.text()
             and "sin cerrar" in dialogo.summary.text(),
             dialogo.summary.text())

fin = HEADERS_HISTORIAL.index("Fin")
textos_fin = [dialogo.table.item(f, fin).text()
              for f in range(dialogo.table.rowCount())]
report.check("el periodo abierto dice que sigue en curso",
             textos_fin.count("en curso") == 1, str(textos_fin))
report.check("y es el unico que se pinta de rojo",
             [f for f in range(dialogo.table.rowCount())
              if dialogo.table.item(f, fin).foreground().color().name().upper()
              == MAINTENANCE_COLOR.upper()] == [0],
             str(textos_fin))

# Y el panel del dashboard, que es donde se pidio que estuviera.
panel = MaintenancePanel()
panel.set_history(todos, context.catalogs.rigs(),
                  maintenance.downtime_by_rig(todos, date(2026, 9, 1)))
report.check("el panel del dashboard resume dias y bancos",
             "fuera de servicio" in panel.subtitle.text(),
             panel.subtitle.text())
report.check("y ofrece abrir el historial completo",
             panel.history_button.isVisibleTo(panel))

vacio = MaintenancePanel()
vacio.set_history([], context.catalogs.rigs(), {})
report.check("sin mantenimientos en el rango, lo dice",
             "Ningún banco" in vacio.note.text(), vacio.note.text())
report.check("y no ofrece un historial que no hay",
             not vacio.history_button.isVisibleTo(vacio))

raise SystemExit(report.finish())

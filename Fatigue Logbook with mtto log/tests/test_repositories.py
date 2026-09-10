"""Verifica repositorios, validacion, filtros y exportacion.

Cubre los pasos 4, 5, 6, 8, 11 y 13 del plan, sobre la copia migrada.
"""

from harness import (
    Report, database_copy, empty_database, make_context, offscreen,
    qt_app, settle,
)


import sys
from datetime import date
from pathlib import Path


from app.db.connection import Database
from app.db.migrations import migrate
from app.db.repositories import (
    AuditRepository, CatalogRepository, FatigueRepository,
    GenericRepository, RotaryRepository,
)
from app.models import FINISHED, ONGOING
from app.services import filtering, validation
from app.services.catalogs import CatalogService
from app.services.excel_export import export_fatigue_ongoing

COPY = database_copy()

failures = []


def check(label, condition, detail=""):
    status = "OK  " if condition else "FALLO"
    print(f"{status} {label}" + (f"   {detail}" if detail else ""))
    if not condition:
        failures.append(label)


def main():
    db = Database(COPY)
    # La copia se migra igual que al arrancar la app: sin esto le faltan
    # las tablas que agregan las migraciones nuevas.
    migrate(db)
    audit = AuditRepository(db)
    catalogs = CatalogService(CatalogRepository(db))
    catalogs.load()

    fatigue = FatigueRepository(db, audit)
    rotary = RotaryRepository(db, audit)
    torsion = GenericRepository(db, audit, "torsion_tests")
    quasi = GenericRepository(db, audit, "quasi_tests")

    print("=== Lectura de repositorios ===")
    # Contra lo que diga la base, no contra un numero escrito a mano: la app
    # esta en uso y estos totales crecen cada semana.
    def rows(table: str) -> int:
        return db.query_one(f"SELECT COUNT(*) FROM {table}")[0]

    all_fatigue = fatigue.list()
    for label, leidos, tabla in (
        ("fatigue", len(all_fatigue), "fatigue_tests"),
        ("rotary", len(rotary.list()), "rotary_tests"),
        ("torsion", len(torsion.list()), "torsion_tests"),
        ("quasi", len(quasi.list()), "quasi_tests"),
    ):
        esperados = rows(tabla)
        check(f"{label}: el repositorio trae todas las filas",
              leidos == esperados, f"{leidos} de {esperados}")

    ongoing = fatigue.list(ONGOING)
    finished = fatigue.list(FINISHED)
    check("ongoing + finished == total",
          len(ongoing) + len(finished) == len(all_fatigue),
          f"{len(ongoing)} + {len(finished)}")

    print("\n=== Totales contra el SQL de la app anterior ===")
    # Consulta literal de fatigue_cmplt_logbook.calculate_total_cycles()
    row = db.query_one("""
        SELECT COALESCE(SUM(COALESCE(cycles1,0) + COALESCE(cycles2,0) +
        COALESCE(cycles3,0) + COALESCE(cycles4,0) + COALESCE(cycles5,0) +
        COALESCE(cycles6,0) + COALESCE(cycles7,0) + COALESCE(cycles8,0) +
        COALESCE(cycles9,0)), 0), COALESCE(SUM(qty_samples), 0)
        FROM fatigue_tests WHERE test_status = 'Finished'
    """)
    old_cycles, old_samples = row[0], row[1]
    new_cycles = sum(t.total_cycles for t in finished)
    new_samples = sum(t.qty_samples for t in finished)

    check("ciclos coinciden", old_cycles == new_cycles,
          f"viejo {old_cycles:,} / nuevo {new_cycles:,}")
    check("piezas coinciden", old_samples == new_samples,
          f"viejo {old_samples:,} / nuevo {new_samples:,}")

    print("\n=== Ida y vuelta de fechas ===")
    with_dates = [t for t in all_fatigue if t.start_date]
    check("fechas parseadas", len(with_dates) == len(all_fatigue),
          f"{len(with_dates)}/{len(all_fatigue)}")
    sample = all_fatigue[0]
    print(f"     ejemplo: {sample.test_batch}  inicio={sample.start_date}"
          f"  fin={sample.end_date}  ciclos={sample.total_cycles:,}")

    print("\n=== Validacion ===")
    codes = catalogs.codes("fatigue")
    try:
        validation.validate_test_batch("242314STF09", codes)
        check("batch valido aceptado", True)
    except validation.ValidationError as e:
        check("batch valido aceptado", False, str(e))

    for bad, why in [("242314XXX09", "clave inexistente"),
                     ("24231STF09", "faltan digitos"),
                     ("", "vacio")]:
        try:
            validation.validate_test_batch(bad, codes)
            check(f"batch invalido rechazado ({why})", False)
        except validation.ValidationError:
            check(f"batch invalido rechazado ({why})", True)

    existing = all_fatigue[0].test_batch
    check("duplicado detectado", fatigue.batch_exists(existing))
    check("duplicado se ignora al editarse a si mismo",
          not fatigue.batch_exists(existing, exclude_id=all_fatigue[0].id))

    print("\n=== Filtros ===")
    total = len(all_fatigue)
    f_customer = filtering.TestFilters(customer=all_fatigue[0].customer)
    by_customer = filtering.apply(all_fatigue, f_customer)
    sql_count = db.query_one(
        "SELECT COUNT(*) AS n FROM fatigue_tests WHERE customer = ?",
        (all_fatigue[0].customer,))["n"]
    check("filtro por cliente cuadra con SQL",
          len(by_customer) == sql_count,
          f"filtro {len(by_customer)} / SQL {sql_count}")

    f_range = filtering.TestFilters(
        date_from=date(2026, 1, 1), date_to=date(2026, 12, 31))
    in_range = filtering.apply(all_fatigue, f_range)
    check("filtro por rango de fechas devuelve un subconjunto",
          0 < len(in_range) < total, f"{len(in_range)} de {total}")

    f_search = filtering.TestFilters(search="STF")
    check("busqueda libre encuentra resultados",
          len(filtering.apply(all_fatigue, f_search)) > 0)

    check("filtro vacio devuelve todo",
          len(filtering.apply(all_fatigue, filtering.TestFilters())) == total)

    print("\n=== Exportacion a Excel ===")
    out = Path(COPY).parent / "reporte.xlsx"
    export_fatigue_ongoing(ongoing, catalogs.colors(), out)
    check("archivo generado", out.is_file(), f"{out.stat().st_size:,} bytes")

    from openpyxl import load_workbook
    wb = load_workbook(out)
    ws = wb.active
    # fila 1 titulo, fila 2 encabezados, N filas, 1 de totales, y si hay
    # piezas suspendidas una en blanco y la nota que explica el naranja.
    suspendidas = sum(1 for t in ongoing for s in t.samples if s.is_suspended)
    esperadas = len(ongoing) + 3 + (2 if suspendidas else 0)
    check("filas correctas en el Excel",
          ws.max_row == esperadas,
          f"{ws.max_row} filas para {len(ongoing)} registros"
          f" y {suspendidas} piezas suspendidas")

    colored = sum(
        1 for row in ws.iter_rows(min_row=3, max_row=ws.max_row - 1)
        for cell in row
        if cell.fill and cell.fill.fgColor.rgb not in (None, "00000000")
    )
    check("hay celdas con color de rig", colored > 0, f"{colored} celdas")
    wb.close()

    # Una pieza suspendida deja la celda de Test Rig vacia, y en una hoja
    # impresa eso es indistinguible de una ranura que nadie uso. Se escribe.
    from app.models import FatigueSample, FatigueTest
    from app.services.excel_export import (
        SUSPENDED_FILL, SUSPENDED_LABEL, SUSPENDED_NOTE,
    )

    con_suspension = FatigueTest(
        test_batch="242314STF09", customer="TOYOTA", start_date=date.today(),
        qty_samples=2,
        samples=[FatigueSample(rig="", cycles=5_000),
                 FatigueSample(rig="I-02-1", cycles=7_000)]
                + [FatigueSample() for _ in range(7)],
    )
    out2 = Path(COPY).parent / "reporte_suspension.xlsx"
    export_fatigue_ongoing([con_suspension], catalogs.colors(), out2)
    wb = load_workbook(out2)
    ws = wb.active
    # Las columnas de muestra se localizan por su nombre: las fijas de la
    # izquierda crecen cada vez que el registro guarda un dato mas.
    cabeceras = [c.value for c in ws[2]]
    rig1 = ws.cell(row=3, column=cabeceras.index("Test Rig 1") + 1)
    rig2 = ws.cell(row=3, column=cabeceras.index("Test Rig 2") + 1)
    check("la celda de la pieza suspendida lo dice",
          rig1.value == SUSPENDED_LABEL, repr(rig1.value))
    check("y va en naranja",
          rig1.fill.fgColor.rgb == SUSPENDED_FILL.fgColor.rgb,
          str(rig1.fill.fgColor.rgb))
    check("la pieza que sigue en banco no se toca",
          rig2.value == "I-02-1", repr(rig2.value))
    textos = [c.value for row in ws.iter_rows() for c in row
              if isinstance(c.value, str)]
    check("la hoja explica el naranja al pie", SUSPENDED_NOTE in textos,
          str([t for t in textos if "Suspendida" in t][:2]))
    wb.close()

    print("\n" + "=" * 62)
    if failures:
        print(f"FALLARON {len(failures)}: " + ", ".join(failures))
        return 1
    print("TODAS LAS COMPROBACIONES PASARON")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())

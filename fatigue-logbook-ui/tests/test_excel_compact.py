"""El reporte de fatigas en curso cabe de lado a lado y no pierde nada.

El de siempre escribe 45 columnas --las nueve ranuras por sus cuatro datos,
esten usadas o no-- y hay que desplazarse para leer un registro. El compacto
baja la muestra a filas: catorce columnas, una fila por pieza.

Lo que se comprueba aqui es que **compacto no signifique incompleto**. Cada
dato que estaba sigue estando, los ciclos se escriben enteros y como numero, y
las piezas de la hoja son exactamente las que la pantalla dibuja -- que es lo
que haria que un reporte firmado dijera una cosa y la bitacora otra.

Nada de contar filas fijas: la base esta en uso y crece. Todo se compara
contra lo que diga la base.
"""

from __future__ import annotations

import tempfile
from pathlib import Path

import harness  # noqa: F401  (deja sys.path listo)
import openpyxl
from harness import Report

from app.models import ONGOING, FatigueSample, FatigueTest
from app.services import maintenance
from app.services.excel_export import (
    MAINTENANCE_FILL,
    MAINTENANCE_LABEL,
    MAINTENANCE_NOTE,
    NO_WO_FILL,
    SUSPENDED_FILL,
    SUSPENDED_LABEL,
    SUSPENDED_NOTE,
)
from app.services.excel_export import build_headers as build_headers_largo
from components.models import FatigueTableModel
from excel_compact import build_headers, export_fatigue_ongoing_compact, used_slots
from theme.rig_palette import RigPalette
from theme.tokens import LIGHT

ESPERADAS = ["ID", "Test Batch", "Cliente", "Requester", "Inicio", "Piezas",
             "Comentarios", "Pieza", "Test Rig", "Resultado", "Ciclos",
             "Modo falla", "Total ciclos", "WO"]

PRIMERA_FILA = 3  # 1 es el titulo y 2 el encabezado


def _export(tests, colores, paused=None):
    """Escribe la hoja en un temporal y la devuelve abierta."""
    destino = Path(tempfile.mkdtemp(prefix="bitacora_excel_")) / "reporte.xlsx"
    ruta = export_fatigue_ongoing_compact(tests, colores, destino,
                                          paused=paused)
    return openpyxl.load_workbook(ruta).active


def _columna(sheet, nombre: str) -> int:
    return [c.value for c in sheet[2]].index(nombre) + 1


def _fila_totales(sheet) -> int:
    """La de TOTALES por su rotulo, no por ser la ultima: debajo van notas."""
    for fila in range(PRIMERA_FILA, sheet.max_row + 1):
        if sheet.cell(row=fila, column=1).value == "TOTALES":
            return fila
    return 0


def _por_prueba(sheet) -> dict[int, list[int]]:
    """Las filas de cada prueba, en orden.

    Se leen arrastrando el ultimo ID visto, que es como las lee una persona:
    las columnas de la prueba estan combinadas sobre sus piezas, asi que
    openpyxl solo devuelve el valor en la fila de arriba de cada combinacion.
    Si la combinacion se rompiera, esto lo notaria.
    """
    filas: dict[int, list[int]] = {}
    actual = None
    for fila in range(PRIMERA_FILA, _fila_totales(sheet)):
        valor = sheet.cell(row=fila, column=1).value
        if isinstance(valor, int):
            actual = valor
        if actual is not None:
            filas.setdefault(actual, []).append(fila)
    return filas


def _texto(valor) -> str:
    """Una celda vacia vuelve de openpyxl como None, no como cadena."""
    return "" if valor is None else str(valor)


def main() -> int:
    report = Report("Reporte compacto de fatigas")
    app = harness.qt_app(LIGHT)
    context = harness.make_context()
    rig_palette = RigPalette(context.database)
    colores = rig_palette.colors()

    registros = context.fatigue.list(ONGOING)
    mantenimientos = context.maintenance.list()
    detenidas = maintenance.paused_slots(mantenimientos)

    report.check("hay pruebas en curso que exportar", bool(registros),
                 f"{len(registros)} pruebas")
    if not registros:
        harness.shutdown(app)
        return report.finish()

    sheet = _export(registros, colores, paused=detenidas)
    filas_por_prueba = _por_prueba(sheet)
    ciclos_col = _columna(sheet, "Ciclos")
    rig_col = _columna(sheet, "Test Rig")

    # --- la hoja cabe de lado a lado -------------------------------------
    report.section("Catorce columnas, con una pieza o con nueve")
    largo = len(build_headers_largo())
    report.check("la hoja compacta tiene 14 columnas",
                 sheet.max_column == len(ESPERADAS),
                 f"{sheet.max_column} columnas, antes {largo}")
    report.check("y son estas, en este orden",
                 [c.value for c in sheet[2]] == ESPERADAS,
                 str([c.value for c in sheet[2]]))
    report.note(f"{largo} - {len(ESPERADAS)} = "
                f"{largo - len(ESPERADAS)} columnas menos")

    # El ancho no puede depender de cuantas piezas traiga lo exportado: ese
    # era justo el problema del reporte anterior.
    una = FatigueTest(test_batch="UNA", customer="X", qty_samples=1,
                      start_date=registros[0].start_date, id=900001)
    una.samples[0] = FatigueSample(rig="I-02-1", result="Falla", cycles=1,
                                   failure_mode="Desgaste")
    nueve = FatigueTest(test_batch="NUEVE", customer="X", qty_samples=9,
                        start_date=registros[0].start_date, id=900002)
    for indice in range(9):
        nueve.samples[indice] = FatigueSample(rig="I-02-1", result="S/Falla",
                                              cycles=1000 + indice)
    for etiqueta, prueba in (("una pieza", una), ("nueve piezas", nueve)):
        hoja = _export([prueba], colores)
        report.check(f"con {etiqueta}, las mismas 14 columnas",
                     hoja.max_column == len(ESPERADAS),
                     f"{hoja.max_column} columnas")

    # --- no se pierde ninguna prueba --------------------------------------
    report.section("Ninguna prueba se queda fuera")
    report.check("cada prueba exportada tiene su fila",
                 set(filas_por_prueba) == {r.id for r in registros},
                 f"{len(filas_por_prueba)} de {len(registros)}")

    vacia = FatigueTest(test_batch="SINPIEZAS", customer="X", qty_samples=3,
                        start_date=registros[0].start_date, id=900003)
    hoja_vacia = _export([vacia], colores)
    report.check("una prueba sin piezas capturadas tambien sale",
                 hoja_vacia.cell(row=PRIMERA_FILA, column=1).value == 900003,
                 str(hoja_vacia.cell(row=PRIMERA_FILA, column=1).value))
    report.check("y sus columnas de pieza van en blanco",
                 all(hoja_vacia.cell(row=PRIMERA_FILA, column=c).value
                     in (None, "") for c in range(8, 13)),
                 str([hoja_vacia.cell(row=PRIMERA_FILA, column=c).value
                      for c in range(8, 13)]))

    # --- una fila por pieza ------------------------------------------------
    report.section("Una fila por pieza usada, ni una mas")
    sobran = []
    for record in registros:
        esperadas = len(used_slots(record)) or 1
        if len(filas_por_prueba.get(record.id, [])) != esperadas:
            sobran.append(f"#{record.id}: "
                          f"{len(filas_por_prueba.get(record.id, []))} filas "
                          f"para {esperadas} piezas")
    report.check("las filas de cada prueba son sus piezas usadas", not sobran,
                 "; ".join(sobran[:3]) if sobran
                 else f"{sum(len(f) for f in filas_por_prueba.values())} filas "
                      f"para {len(registros)} pruebas")

    report.section("Las dos vistas coinciden")
    # El reporte y la tabla preguntan lo mismo (``models.has_content``) pero
    # por caminos distintos. Si divergen, el papel dice una cosa y la pantalla
    # otra, y no hay forma de saber cual creer.
    modelo = FatigueTableModel(rig_palette, show_end_date=False, compact=True)
    modelo.set_maintenance(
        maintenance.stopped_by_test(context.fatigue.list(), mantenimientos),
        detenidas)
    modelo.set_records(registros)
    distintas = []
    for record in registros:
        chips = len(modelo.chips(record))
        filas = len(filas_por_prueba.get(record.id, []))
        if chips and chips != filas:
            distintas.append(f"#{record.id}: {chips} chips, {filas} filas")
    report.check("la hoja ensenia las mismas piezas que la tabla",
                 not distintas,
                 "; ".join(distintas[:3]) if distintas else "coinciden todas")

    # --- los ciclos, completos --------------------------------------------
    report.section("Los ciclos van completos")
    abreviados = []
    no_numericos = []
    for filas in filas_por_prueba.values():
        for fila in filas:
            celda = sheet.cell(row=fila, column=ciclos_col)
            if celda.value not in (None, "") and not isinstance(celda.value,
                                                                int):
                no_numericos.append(f"fila {fila}: {celda.value!r}")
            if any(letra in _texto(celda.value).upper() for letra in "KM"):
                abreviados.append(f"fila {fila}: {celda.value!r}")
    report.check("ningun ciclo esta abreviado", not abreviados,
                 "; ".join(abreviados[:3]) if abreviados
                 else "sin K ni M en toda la columna")
    report.check("los ciclos son numero, no texto", not no_numericos,
                 "; ".join(no_numericos[:3]) if no_numericos
                 else "toda la columna es entera")
    report.check("y llevan separador de miles",
                 sheet.cell(row=PRIMERA_FILA,
                            column=ciclos_col).number_format == "#,##0",
                 sheet.cell(row=PRIMERA_FILA,
                            column=ciclos_col).number_format)

    # El caso que motiva la regla: el numero mas grande de la base, escrito
    # tal cual. Abreviarlo lo dejaria en '8.1M' y nadie podria recapturarlo.
    mayor = max((s.cycles or 0) for r in registros for s in r.samples)
    escritos = {sheet.cell(row=f, column=ciclos_col).value
                for filas in filas_por_prueba.values() for f in filas}
    report.check("el ciclo mas alto de la base esta escrito entero",
                 mayor in escritos, f"{mayor:,}")

    # --- nada mas se pierde ------------------------------------------------
    report.section("Cada dato de cada pieza sigue ahi")
    perdidos = []
    comprobados = 0
    for record in registros:
        for fila, numero in zip(filas_por_prueba[record.id],
                                used_slots(record)):
            sample = record.samples[numero - 1]
            leido = {
                "Pieza": sheet.cell(row=fila, column=8).value,
                "Resultado": _texto(sheet.cell(row=fila, column=10).value),
                "Ciclos": sheet.cell(row=fila, column=ciclos_col).value,
                "Modo falla": _texto(sheet.cell(row=fila, column=12).value),
            }
            esperado = {
                "Pieza": numero,
                "Resultado": sample.result or "",
                "Ciclos": sample.cycles if sample.cycles is not None else None,
                "Modo falla": sample.failure_mode or "",
            }
            if leido["Ciclos"] == "":
                leido["Ciclos"] = None
            for campo, valor in esperado.items():
                comprobados += 1
                if leido[campo] != valor:
                    perdidos.append(f"#{record.id} pieza {numero} {campo}: "
                                    f"{leido[campo]!r} != {valor!r}")
    report.check("resultado, ciclos y modo de falla, dato por dato",
                 not perdidos,
                 "; ".join(perdidos[:3]) if perdidos
                 else f"{comprobados} datos comprobados")

    # --- las marcas siguen ahi ---------------------------------------------
    report.section("Las marcas de color dicen lo mismo")
    marcadas = {"detenida": 0, "suspendida": 0, "con banco": 0}
    mal = []
    for record in registros:
        for fila, numero in zip(filas_por_prueba[record.id],
                                used_slots(record)):
            sample = record.samples[numero - 1]
            celda = sheet.cell(row=fila, column=rig_col)
            valor = _texto(celda.value)
            if (record.id, numero) in detenidas:
                marcadas["detenida"] += 1
                if (valor != MAINTENANCE_LABEL
                        or celda.fill.fgColor.rgb
                        != MAINTENANCE_FILL.fgColor.rgb):
                    mal.append(f"#{record.id} pieza {numero} detenida: "
                               f"{valor!r} {celda.fill.fgColor.rgb}")
            elif not record.is_finished and sample.is_suspended:
                marcadas["suspendida"] += 1
                if (valor != SUSPENDED_LABEL
                        or celda.fill.fgColor.rgb
                        != SUSPENDED_FILL.fgColor.rgb):
                    mal.append(f"#{record.id} pieza {numero} suspendida: "
                               f"{valor!r} {celda.fill.fgColor.rgb}")
            elif colores.get(valor):
                marcadas["con banco"] += 1
                esperado = colores[valor].lstrip("#").upper()
                if not _texto(celda.fill.fgColor.rgb).upper().endswith(
                        esperado):
                    mal.append(f"#{record.id} pieza {numero} banco {valor}: "
                               f"{celda.fill.fgColor.rgb} != {esperado}")
    report.check("detenida, suspendida y banco, cada una con su color",
                 not mal, "; ".join(mal[:3]) if mal else str(marcadas))

    report.check("hay piezas detenidas que comprobar",
                 marcadas["detenida"] > 0, f"{marcadas['detenida']} piezas")
    report.check("y piezas con banco de verdad",
                 marcadas["con banco"] > 0, f"{marcadas['con banco']} piezas")

    # El resaltado de "sin Work Order" cubre la prueba entera, no solo su
    # primera fila: media prueba resaltada se lee como si la otra media si la
    # tuviera.
    sin_wo = [r for r in registros if not r.wo_status]
    if sin_wo:
        record = sin_wo[0]
        filas = filas_por_prueba[record.id]
        report.check("una prueba sin WO va resaltada en todas sus filas",
                     all(sheet.cell(row=f, column=1).fill.fgColor.rgb
                         == NO_WO_FILL.fgColor.rgb for f in filas),
                     f"#{record.id}, {len(filas)} filas")
    else:
        report.note("no hay pruebas sin Work Order en curso; nada que medir")

    textos = {_texto(c.value) for fila in sheet.iter_rows(min_row=PRIMERA_FILA)
              for c in fila}
    if marcadas["suspendida"]:
        report.check("la nota explica el naranja", SUSPENDED_NOTE in textos,
                     "presente")
    if marcadas["detenida"]:
        report.check("y la nota explica el rojo", MAINTENANCE_NOTE in textos,
                     "presente")

    # --- combinadas --------------------------------------------------------
    report.section("Los datos de la prueba se combinan sobre sus piezas")
    varias = [r for r in registros if len(filas_por_prueba[r.id]) > 1]
    report.check("hay pruebas de mas de una pieza", bool(varias),
                 f"{len(varias)} pruebas")
    if varias:
        record = varias[0]
        filas = filas_por_prueba[record.id]
        rangos = {(r.min_row, r.max_row, r.min_col)
                  for r in sheet.merged_cells.ranges}
        report.check("el ID abarca las filas de sus piezas",
                     (filas[0], filas[-1], 1) in rangos,
                     f"#{record.id}: filas {filas[0]}-{filas[-1]}")
        report.check("y el total de ciclos tambien",
                     (filas[0], filas[-1], 13) in rangos,
                     f"filas {filas[0]}-{filas[-1]}")

    # --- totales -----------------------------------------------------------
    report.section("Los totales cuadran")
    totales = _fila_totales(sheet)
    report.check("hay fila de TOTALES", totales > 0, f"fila {totales}")
    esperado_ciclos = sum(r.total_cycles for r in registros)
    report.check("el total de ciclos es el de los registros",
                 sheet.cell(row=totales, column=13).value == esperado_ciclos,
                 f"{sheet.cell(row=totales, column=13).value:,} vs "
                 f"{esperado_ciclos:,}")
    # El mismo numero por otro camino: sumando las piezas una a una. Si
    # difieren, alguna pieza no llego a la hoja.
    report.check("y sumando pieza a pieza da lo mismo",
                 sheet.cell(row=totales, column=ciclos_col).value
                 == esperado_ciclos,
                 f"{sheet.cell(row=totales, column=ciclos_col).value:,}")
    report.check("las piezas declaradas se suman igual",
                 sheet.cell(row=totales, column=6).value
                 == sum(r.qty_samples for r in registros),
                 str(sheet.cell(row=totales, column=6).value))

    report.section("La hoja se arma con los ayudantes del original")
    report.check("encabezado congelado", sheet.freeze_panes == "D3",
                 str(sheet.freeze_panes))
    report.check("autofiltro sobre las 14 columnas",
                 _texto(sheet.auto_filter.ref).startswith("A2:N"),
                 str(sheet.auto_filter.ref))
    report.check("la columna del banco cabe 'Mantenimiento'",
                 sheet.column_dimensions["I"].width >= 14,
                 f"{sheet.column_dimensions['I'].width} de ancho")
    report.check("las cabeceras las da build_headers",
                 build_headers() == ESPERADAS, str(build_headers()))

    harness.shutdown(app)
    return report.finish()


if __name__ == "__main__":
    raise SystemExit(main())
